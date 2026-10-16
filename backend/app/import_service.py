from openpyxl import load_workbook
from openpyxl import load_workbook
from .models import Measurement

def inspect_excel(path: str):

    workbook = load_workbook(path)

    sheet = workbook.active

    headers = [cell.value for cell in sheet[1]]

    return headers




def read_excel_rows(path: str):

    workbook = load_workbook(path)

    sheet = workbook.active

    rows = []

    for row in sheet.iter_rows(min_row=2, values_only=True):

        rows.append(
            {
                "timestamp": row[0],
                "systolic": row[1],
                "diastolic": row[2],
                "pulse": row[3],
                "note": row[4],
                "person": row[5],
            }
        )

    return rows


def import_measurements(path: str, db):

    workbook = load_workbook(path)

    sheet = workbook.active

    imported = 0
    skipped = 0

    for row in sheet.iter_rows(min_row=2, values_only=True):

        try:

            measurement = Measurement(
                measurement_time=row[0],
                systolic=row[1],
                diastolic=row[2],
                pulse=row[3],
                note=row[4] if row[4] else "",
                person=row[5]
            )

            db.add(measurement)

            imported += 1

        except Exception:
            skipped += 1

    db.commit()

    return {
        "imported": imported,
        "skipped": skipped
    }